"""
The statement viewer.

`/reports` is the financial pack: the OHADA statutory statements, the IFRS
presentation of the same ledger, and the bridge that reconciles the two.
`/asset-reports` is the register, depreciation, compliance and cost-of-ownership
side. Both export to Excel and PDF.
"""

from __future__ import annotations

import io
from datetime import datetime

from nicegui import ui

from .. import api
from ..i18n import bilingual, framework as current_framework, set_framework, t
from ..ui import chrome, widgets as w
from ..ui.shell import page

OHADA_SLUGS = ('bilan', 'resultat', 'tafire')
IFRS_SLUGS = ('sofp', 'pl', 'oci', 'cash-flow', 'equity', 'ppe-note')


def _ohada_tabs() -> dict[str, str]:
    return {'bilan': t('st.bilan'), 'resultat': t('st.resultat'), 'tafire': t('st.tafire')}


def _ifrs_tabs() -> dict[str, str]:
    return {
        'sofp': t('st.sofp'), 'pl': t('st.pl'), 'oci': t('st.oci'),
        'cash-flow': t('st.cf'), 'equity': t('st.soce'), 'ppe-note': t('st.ppe'),
    }


@page('/reports', title='Financial statements', permissions=('report.financial',))
async def reports_page(statement: str = '') -> None:
    currency = api.user().get('functional_currency', 'XAF')
    catalog = await api.safe(api.get('/reports/catalog'))
    if not catalog:
        return
    years = catalog.get('fiscal_years') or []
    current = next((y for y in years if y.get('status') != 'closed'), years[0] if years else {})

    # The header OHADA / IFRS toggle is the source of truth. A statement slug
    # only picks a tab inside that framework; `bridge` is the exception.
    header_fw = current_framework()
    if statement == 'bridge':
        framework_of = 'bridge'
        default_statement = 'sofp' if header_fw == 'ifrs' else 'bilan'
    elif header_fw == 'ifrs':
        framework_of = 'ifrs'
        default_statement = statement if statement in IFRS_SLUGS else 'sofp'
    elif statement in OHADA_SLUGS:
        framework_of = 'ohada'
        default_statement = statement
    else:
        framework_of = 'ohada'
        default_statement = 'bilan'
    state = {'from': str(current.get('start_date', ''))[:10],
             'to': str(current.get('end_date', ''))[:10],
             'framework': framework_of,
             'statement': default_statement}

    org = catalog.get('organisation') or {}
    w.page_header(
        t('st.title'),
        t('st.sub', org=org.get('name', ''),
          tier=str(org.get('reporting_tier', '')).capitalize(),
          ccy=org.get('functional_currency', '')),
        emoji=chrome.PAGE['st.title'],
    )

    export_state: dict = {'rows': [], 'title': '', 'headers': []}

    async def reload() -> None:
        container.clear()
        with container:
            if state['framework'] == 'ohada':
                await _ohada(state, currency, export_state)
            elif state['framework'] == 'ifrs':
                await _ifrs(state, currency, export_state)
            else:
                await _bridge(state, currency)

    with ui.element('div').classes('lc-card p-4 w-full lc-no-print'):
        with w.toolbar():
            framework = ui.toggle({'ohada': t('st.ohada'), 'ifrs': t('st.ifrs'),
                                   'bridge': t('st.bridge')}, value=state['framework']) \
                .props('unelevated dense no-caps').classes('rounded-lg')
            framework.bind_value(state, 'framework')
            date_from = ui.input(t('common.from'), value=state['from']).props('outlined dense type=date') \
                .classes('w-40')
            date_from.bind_value(state, 'from')
            date_to = ui.input(t('common.to'), value=state['to']).props('outlined dense type=date').classes('w-40')
            date_to.bind_value(state, 'to')
            ui.button(t('common.print'), icon='print',
                      on_click=lambda: ui.run_javascript('window.print()')).props('outline dense')
            ui.button(t('common.excel'), icon='table_view',
                      on_click=lambda: _export_excel(export_state, currency)).props('outline dense')
            ui.button(t('common.pdf'), icon='picture_as_pdf',
                      on_click=lambda: _export_pdf(export_state, currency, org)).props('outline dense')

        statement_row = ui.row().classes('w-full pt-2')

        def draw_statement_picker() -> None:
            statement_row.clear()
            options = (_ohada_tabs() if state['framework'] == 'ohada'
                       else _ifrs_tabs() if state['framework'] == 'ifrs' else None)
            if not options:
                return
            if state['statement'] not in options:
                state['statement'] = next(iter(options))
            with statement_row:
                picker = ui.toggle(options, value=state['statement']) \
                    .props('flat dense no-caps').classes('text-sm')
                picker.bind_value(state, 'statement')
                picker.on_value_change(reload)

        async def switch_framework() -> None:
            if state['framework'] in ('ohada', 'ifrs'):
                set_framework(state['framework'])
                if state['framework'] != header_fw:
                    ui.navigate.reload()
                    return
            draw_statement_picker()
            await reload()

        framework.on_value_change(switch_framework)
        for element in (date_from, date_to):
            element.on_value_change(reload)
        draw_statement_picker()

    container = ui.column().classes('w-full gap-5')
    await reload()


# --- OHADA ------------------------------------------------------------------


async def _ohada(state: dict, currency: str, export_state: dict) -> None:
    slug = state['statement']
    data = await api.safe(api.get(f'/reports/ohada/{slug}',
                                  **{'from': state['from'], 'to': state['to']}))
    if not data:
        return

    check = data.get('check') or {}
    if slug == 'resultat':
        w.kpi_row()
        with w.kpi_row():
            w.kpi(t('st.net_period'), w.money(check.get('net_result'), currency),
                  hint=t('st.net_result_hint'), icon='trending_up'
                  if float(check.get('net_result') or 0) >= 0 else 'trending_down',
                  tone='teal' if float(check.get('net_result') or 0) >= 0 else 'rose')
    elif check.get('balanced') is not None:
        w.check_banner([check], currency)

    lines = data.get('lines') or []
    is_bilan = slug == 'bilan'
    headers = (['', t('st.gross', ccy=currency), t('st.depr', ccy=currency), t('st.net', ccy=currency)]
               if is_bilan else ['', f'{currency}'])

    # The Bilan is presented as two faces; splitting on `side` keeps ACTIF and
    # PASSIF readable rather than running them into one column.
    if is_bilan:
        for side, title in (('assets', t('st.actif')),
                            ('liabilities', t('st.passif'))):
            face = [line for line in lines if line.get('side') == side]
            if not face:
                continue
            body = w.section(title, f"{w.as_date(state['from'])} to {w.as_date(state['to'])}")
            with body:
                w.statement_table(headers, [_ohada_row(line, True) for line in face], currency)
    else:
        body = w.section(bilingual(data, fr_keys=('title',), en_keys=('title_en', 'title')),
                         f"{w.as_date(state['from'])} {t('common.to').lower()} {w.as_date(state['to'])}")
        with body:
            w.statement_table(headers, [_ohada_row(line, False) for line in lines], currency)

    export_state.update(
        title=f"{data.get('title', slug)} — {w.as_date(state['from'])} to {w.as_date(state['to'])}",
        headers=['Code', 'Label'] + headers[1:],
        rows=[[line.get('line_code'), line.get('label')]
              + ([line.get('gross'), line.get('depreciation'), line.get('net')] if is_bilan
                 else [line.get('net')])
              for line in lines],
    )


def _ohada_row(line: dict, with_columns: bool) -> dict:
    kind = ('total' if line.get('line_type') == 'total'
            else 'subtotal' if line.get('line_type') in ('subtotal', 'total_partial')
            else 'header' if line.get('line_type') == 'header'
            else 'detail')
    if line.get('is_bold') and kind == 'detail':
        kind = 'subtotal'
    values = ([line.get('gross'), line.get('depreciation'), line.get('net')] if with_columns
              else [line.get('net')])
    return {'code': line.get('line_code'), 'label': bilingual(line),
            'note': line.get('note_ref'), 'kind': kind, 'values': values}


# --- IFRS -------------------------------------------------------------------


async def _ifrs(state: dict, currency: str, export_state: dict) -> None:
    slug = state['statement']
    data = await api.safe(api.get(f'/reports/ifrs/{slug}',
                                  **{'from': state['from'], 'to': state['to']}))
    if not data:
        return

    if slug == 'ppe-note':
        _ppe_note(data, currency, state, export_state)
        return
    if slug == 'equity':
        _equity(data, currency, state, export_state)
        return

    check = data.get('check')
    if check and check.get('balanced') is not None:
        w.check_banner([check], currency)

    totals = data.get('totals') or {}
    if slug == 'sofp':
        with w.kpi_row():
            w.kpi(t('st.total_assets'), w.money(totals.get('total_assets'), currency),
                  hint=t('st.hint_nc_c'), icon='account_balance', tone='sky')
            w.kpi(t('st.total_equity'), w.money(totals.get('total_equity'), currency),
                  hint=t('st.hint_incl_result'), icon='savings', tone='teal')
            w.kpi(t('st.total_liabilities'), w.money(totals.get('total_liabilities'), currency),
                  hint=t('st.hint_nc_c'), icon='credit_card', tone='violet')
    elif slug == 'pl':
        with w.kpi_row():
            w.kpi(t('st.total_income'), w.money(totals.get('total_income'), currency),
                  hint=t('st.hint_revenue'), icon='payments', tone='sky')
            w.kpi(t('st.operating_profit'), w.money(totals.get('operating_profit'), currency),
                  hint=t('st.hint_before_fin'), icon='query_stats', tone='teal')
            w.kpi(t('st.profit_year'), w.money(totals.get('profit_for_the_year'), currency),
                  hint=t('st.hint_after_fin'), icon='trending_up',
                  tone='teal' if float(totals.get('profit_for_the_year') or 0) >= 0 else 'rose')
    elif slug == 'oci':
        with w.kpi_row():
            w.kpi(t('st.total_oci'), w.money(totals.get('total_oci'), currency),
                  hint=t('st.hint_reval'),
                  icon='stacked_line_chart', tone='violet')

    rows: list[dict] = []
    if slug == 'cash-flow':
        for section in data.get('sections') or []:
            heading = bilingual(section)
            rows.append({'label': heading, 'kind': 'header', 'values': []})
            for line in section.get('lines') or []:
                rows.append({'label': bilingual(line), 'values': [line.get('amount')]})
            rows.append({'label': t('st.net_cash_from', name=heading),
                         'kind': 'subtotal', 'values': [section.get('subtotal') or 0]})
        rows += [
            {'label': t('st.net_change_cash'), 'kind': 'subtotal',
             'values': [totals.get('net_change')]},
            {'label': t('st.opening_cash'),
             'values': [totals.get('opening_cash')]},
            {'label': t('st.closing_cash'), 'kind': 'total',
             'values': [totals.get('closing_cash_calculated')]},
        ]
    else:
        for section in data.get('sections') or []:
            heading = bilingual(section)
            rows.append({'label': heading, 'kind': 'header', 'values': []})
            for line in section.get('lines') or []:
                rows.append({'code': line.get('code'), 'label': bilingual(line),
                             'note': line.get('note_ref'), 'values': [line.get('amount')]})
            # A subtotal of nil is still a figure, so show it rather than a gap.
            rows.append({'label': t('st.total_of', name=heading), 'kind': 'subtotal',
                         'values': [section.get('subtotal') or 0]})
        if slug == 'sofp':
            rows += [
                {'label': t('st.total_assets'), 'kind': 'total', 'values': [totals.get('total_assets')]},
                {'label': t('st.total_eq_liab'), 'kind': 'total',
                 'values': [totals.get('total_equity_and_liabilities')]},
            ]
        elif slug == 'pl':
            rows += [
                {'label': t('st.operating_profit'), 'kind': 'subtotal',
                 'values': [totals.get('operating_profit')]},
                {'label': t('st.profit_before_tax'), 'kind': 'subtotal',
                 'values': [totals.get('profit_before_tax')]},
                {'label': t('st.profit_year'), 'kind': 'total',
                 'values': [totals.get('profit_for_the_year')]},
            ]
        elif slug == 'oci':
            rows.append({'label': t('st.oci_year'), 'kind': 'total',
                         'values': [totals.get('total_oci')]})

    body = w.section(_ifrs_tabs().get(slug, slug),
                     f"{w.as_date(state['from'])} {t('common.to').lower()} {w.as_date(state['to'])} · "
                     f"{t('st.same_books')}")
    with body:
        w.statement_table(['', currency], rows, currency)

    export_state.update(
        title=f"{_ifrs_tabs().get(slug, slug)} — {w.as_date(state['from'])} to {w.as_date(state['to'])}",
        headers=['Label', currency],
        rows=[[row['label'], (row['values'][0] if row.get('values') else None)] for row in rows],
    )


def _equity(data: dict, currency: str, state: dict, export_state: dict) -> None:
    check = data.get('check')
    if check:
        w.check_banner([check], currency)
    columns = data.get('columns') or []
    rows = data.get('rows') or []

    body = w.section(t('st.soce'),
                     f"{w.as_date(state['from'])} {t('common.to').lower()} {w.as_date(state['to'])}")
    with body:
        w.statement_table(
            [''] + [bilingual(c) for c in columns] + [t('common.total')],
            [{'label': bilingual(row),
              'kind': 'total' if row.get('is_total') else 'detail',
              'values': [row['values'].get(c['key']) for c in columns] + [row.get('total')]}
             for row in rows],
            currency,
        )

    export_state.update(
        title=f"{t('st.soce')} — {w.as_date(state['from'])} to {w.as_date(state['to'])}",
        headers=['Label'] + [bilingual(c) for c in columns] + [t('common.total')],
        rows=[[bilingual(row)] + [row['values'].get(c['key']) for c in columns] + [row.get('total')]
              for row in rows],
    )


def _ppe_note(data: dict, currency: str, state: dict, export_state: dict) -> None:
    checks = data.get('checks') or []
    w.check_banner(checks, currency)
    totals = data.get('totals') or {}
    classes = data.get('classes') or []

    with w.kpi_row():
        w.kpi(t('st.cost_end'), w.money(totals.get('closing_cost'), currency),
              hint=t('st.cost_end_hint', amount=w.money(totals.get('additions'), currency)),
              icon='request_quote', tone='sky')
        w.kpi(t('st.accum_depr'), w.money(totals.get('closing_accum_depr'), currency),
              hint=t('st.accum_depr_hint', amount=w.money(totals.get('depreciation_charge'), currency)),
              icon='trending_down', tone='violet')
        w.kpi(t('st.accum_imp'), w.money(totals.get('accum_impairment'), currency),
              hint=t('st.accum_imp_hint'), icon='south', tone='rose')
        w.kpi(t('st.carrying'), w.money(totals.get('carrying_amount'), currency),
              hint=t('st.carrying_hint'),
              icon='account_balance', tone='teal')

    body = w.section(t('st.ppe_title'),
                     t('st.ppe_sub', basis=data.get('basis', 'ifrs').upper()))
    with body:
        headers = [''] + [bilingual(c) for c in classes] + [t('common.total')]
        spec = [
            (t('st.ppe.opening_cost'), 'opening_cost', 'detail'),
            (t('st.ppe.additions'), 'additions', 'detail'),
            (t('st.ppe.disposals'), 'disposals_cost', 'detail'),
            (t('st.ppe.revaluation'), 'revaluation', 'detail'),
            (t('st.ppe.closing_cost'), 'closing_cost', 'subtotal'),
            (t('st.ppe.opening_accum'), 'opening_accum_depr', 'detail'),
            (t('st.ppe.charge'), 'depreciation_charge', 'detail'),
            (t('st.ppe.elim'), 'disposals_accum_depr', 'detail'),
            (t('st.ppe.closing_accum'), 'closing_accum_depr', 'subtotal'),
            (t('st.ppe.impairment'), 'accum_impairment', 'detail'),
            (t('st.ppe.carrying_end'), 'carrying_amount', 'total'),
        ]
        rows = [{'label': label, 'kind': kind,
                 'values': [c.get(key) for c in classes] + [totals.get(key)]}
                for label, key, kind in spec]
        w.statement_table(headers, rows, currency)

    export_state.update(
        title=f"PP&E note — {w.as_date(state['from'])} to {w.as_date(state['to'])}",
        headers=headers,
        rows=[[row['label']] + row['values'] for row in rows],
    )


# --- the bridge -------------------------------------------------------------


async def _bridge(state: dict, currency: str) -> None:
    data = await api.safe(api.get('/reports/dual-standard-bridge',
                                  **{'from': state['from'], 'to': state['to']}))
    if not data:
        return

    reconciled = bool(data.get('reconciled'))
    w.check_banner([{'label': t('st.bridge_check'),
                     'balanced': reconciled, 'difference': data.get('difference')}], currency)

    with w.kpi_row():
        w.kpi(t('st.ohada_result'), w.money(data.get('ohada_result'), currency),
              hint=t('st.ohada_result_hint'), icon='balance', tone='sky')
        w.kpi(t('st.ifrs_adj'), w.money(data.get('ifrs_only_adjustments'), currency),
              hint=t('st.ifrs_adj_hint'),
              icon='call_split', tone='amber')
        w.kpi(t('st.ifrs_result'), w.money(data.get('ifrs_result'), currency),
              hint=t('st.ifrs_result_hint'), icon='trending_up', tone='violet')

    left, right = w.two_column()
    with left:
        body = w.section(t('st.how_reconcile'), t('st.how_reconcile_sub'))
        with body:
            w.statement_table(
                ['', currency],
                [
                    {'label': t('st.result_ohada'),
                     'values': [data.get('ohada_result')]},
                    {'label': t('st.adj_ifrs'),
                     'values': [data.get('ifrs_only_adjustments')]},
                    {'label': t('st.result_ifrs'), 'kind': 'total',
                     'values': [data.get('ifrs_result')]},
                ],
                currency,
            )

        adjustments = data.get('adjustments') or []
        body = w.section(t('st.ifrs_only_n', n=w.number(len(adjustments))),
                         t('st.ifrs_only_sub'))
        with body:
            if not adjustments:
                w.empty_state(t('st.no_ifrs_adj'), icon='call_split',
                              hint=t('st.same_result'))
            else:
                w.table([w.column('entry_no', t('je.entry_no'), classes='lc-num'),
                         w.date_column('entry_date', t('je.date')),
                         w.column('description', t('je.description')),
                         w.column('source', 'Source'),
                         w.money_column('profit_effect', f"{t('st.ifrs_result')} ({currency})",
                                        currency=currency)],
                        adjustments, pagination=20)

    with right:
        body = w.section(t('st.by_framework'), t('st.by_framework_sub'))
        with body:
            rows = data.get('entries_by_framework') or []
            labels = {'both': t('st.fw_both'), 'ohada': t('st.fw_ohada'), 'ifrs': t('st.fw_ifrs')}
            for row in rows:
                with ui.row().classes('w-full items-center justify-between py-2 no-wrap '
                                      'border-b border-slate-100 last:border-0'):
                    with ui.column().classes('gap-0'):
                        ui.label(labels.get(row['framework'], row['framework'])
                                 ).classes('text-xs font-medium text-slate-700')
                        ui.label(t('common.entries', n=w.number(row['entries']))
                                 ).classes('text-[11px] text-slate-500')
                    ui.label(w.money(row['total'], currency)).classes('lc-num text-sm font-semibold')
        w.note(t('st.bridge_note'), icon='balance')


# --- asset reports ----------------------------------------------------------


@page('/asset-reports', title='Asset and compliance reports', permissions=('report.asset',))
async def asset_reports_page() -> None:
    currency = api.user().get('functional_currency', 'XAF')
    catalog = await api.safe(api.get('/reports/catalog'), notify=False) or {}
    years = catalog.get('fiscal_years') or []
    current = next((y for y in years if y.get('status') != 'closed'), years[0] if years else {})

    state = {'from': str(current.get('start_date', ''))[:10],
             'to': str(current.get('end_date', ''))[:10],
             'asOf': str(current.get('end_date', ''))[:10],
             'basis': 'ohada', 'groupBy': 'category', 'report': 'register'}
    export_state: dict = {'rows': [], 'title': '', 'headers': []}

    w.page_header(t('arep.title'), t('arep.sub'), emoji=chrome.PAGE['arep.title'])

    async def reload() -> None:
        container.clear()
        with container:
            report = state['report']
            if report == 'register':
                await _register_report(state, currency, export_state)
            elif report == 'reconciliation':
                await _reconciliation(state, currency, export_state)
            elif report == 'depreciation':
                await _depreciation_report(state, currency, export_state)
            elif report == 'compliance':
                await _compliance_report(state, export_state)
            else:
                await _tco_report(state, currency, export_state)

    with ui.element('div').classes('lc-card p-4 w-full lc-no-print'):
        with w.toolbar():
            report = ui.toggle({'register': t('arep.register'), 'reconciliation': t('arep.recon'),
                                'depreciation': t('arep.depr'), 'compliance': t('arep.compliance'),
                                'tco': t('arep.tco')}, value='register') \
                .props('unelevated dense no-caps').classes('rounded-lg')
            report.bind_value(state, 'report')
            date_from = ui.input(t('common.from'), value=state['from']).props('outlined dense type=date') \
                .classes('w-40')
            date_from.bind_value(state, 'from')
            date_to = ui.input(t('common.to'), value=state['to']).props('outlined dense type=date').classes('w-40')
            date_to.bind_value(state, 'to')
            as_of = ui.input(t('common.as_at'), value=state['asOf']).props('outlined dense type=date') \
                .classes('w-40')
            as_of.bind_value(state, 'asOf')
            basis = ui.select({'ohada': t('arep.ohada_basis'), 'ifrs': t('arep.ifrs_basis')}, value='ohada',
                              label=t('arep.basis')).props('outlined dense options-dense').classes('w-40')
            basis.bind_value(state, 'basis')
            group = ui.select({'category': t('arep.by_category'), 'domain': t('arep.by_domain'),
                               'cost_center': t('arep.by_cc'), 'location': t('arep.by_loc'),
                               'funding_source': t('arep.by_funding')},
                              value='category', label=t('arep.group')) \
                .props('outlined dense options-dense').classes('w-44')
            group.bind_value(state, 'groupBy')
            ui.button(t('common.print'), icon='print',
                      on_click=lambda: ui.run_javascript('window.print()')).props('outline dense')
            ui.button(t('common.excel'), icon='table_view',
                      on_click=lambda: _export_excel(export_state, currency)).props('outline dense')
            ui.button(t('common.pdf'), icon='picture_as_pdf',
                      on_click=lambda: _export_pdf(export_state, currency,
                                                   catalog.get('organisation') or {})
                      ).props('outline dense')
            for element in (report, date_from, date_to, as_of, basis, group):
                element.on_value_change(reload)

    container = ui.column().classes('w-full gap-5')
    await reload()


async def _register_report(state: dict, currency: str, export_state: dict) -> None:
    data = await api.safe(api.get('/reports/asset-register', asOf=state['asOf'],
                                  basis=state['basis'], groupBy=state['groupBy']))
    if not data:
        return
    totals = data.get('totals') or {}
    rows = data.get('rows') or []

    with w.kpi_row():
        w.kpi('Assets on register', w.number(totals.get('asset_count')),
              hint=f"Capitalised as at {w.as_date(state['asOf'])}", icon='inventory_2', tone='slate')
        w.kpi('Gross cost', w.money(totals.get('gross'), currency),
              hint=f"{state['basis'].upper()} basis", icon='request_quote', tone='sky')
        w.kpi('Accumulated depreciation', w.money(totals.get('accumulated_depreciation'), currency),
              hint='Charged to date', icon='trending_down', tone='violet')
        w.kpi('Net book value', w.money(totals.get('net'), currency),
              hint='What the statements carry', icon='account_balance', tone='teal')
        w.kpi('Lifetime maintenance', w.money(totals.get('maintenance_cost'), currency),
              hint='Spent keeping these assets running', icon='build', tone='amber')

    body = w.section(f"Asset register as at {w.as_date(state['asOf'])}",
                     f"{state['basis'].upper()} basis, grouped by "
                     f"{state['groupBy'].replace('_', ' ')}")
    with body:
        if not rows:
            w.empty_state('No capitalised asset at that date.', icon='inventory_2')
            return
        headers = ['', 'Count', f'Cost ({currency})', f'Depreciation ({currency})',
                   f'Impairment ({currency})', f'Net book value ({currency})',
                   f'Maintenance ({currency})']
        statement_rows = [
            {'code': row.get('group_code'), 'label': w.humanise(row.get('group_name')),
             'values': [row.get('asset_count'), row.get('gross'),
                        -abs(float(row.get('accumulated_depreciation') or 0)),
                        -abs(float(row.get('impairment') or 0)),
                        row.get('net'), row.get('maintenance_cost')]}
            for row in rows
        ] + [{'label': 'Total', 'kind': 'total',
              'values': [totals.get('asset_count'), totals.get('gross'),
                         -abs(float(totals.get('accumulated_depreciation') or 0)),
                         -abs(float(totals.get('impairment') or 0)),
                         totals.get('net'), totals.get('maintenance_cost')]}]
        w.statement_table(headers, statement_rows, currency)

    export_state.update(title=f"Asset register as at {w.as_date(state['asOf'])}",
                        headers=['Code'] + headers[0:1] + headers[1:],
                        rows=[[r.get('code'), r['label']] + r['values'] for r in statement_rows])


async def _reconciliation(state: dict, currency: str, export_state: dict) -> None:
    data = await api.safe(api.get('/reports/register-reconciliation', asOf=state['asOf'],
                                  basis=state['basis']))
    if not data:
        return
    totals = data.get('totals') or {}
    reconciled = bool(totals.get('reconciled'))

    w.check_banner([
        {'label': 'Register cost less ledger cost', 'balanced': float(totals.get('cost_variance') or 0) == 0,
         'difference': totals.get('cost_variance')},
        {'label': 'Register depreciation less ledger depreciation',
         'balanced': float(totals.get('accum_variance') or 0) == 0,
         'difference': totals.get('accum_variance')},
    ], currency, subject='reconciliation')

    with w.kpi_row():
        w.kpi('Register cost', w.money(totals.get('register_cost'), currency),
              hint='Summed from the asset register', icon='inventory_2', tone='sky')
        w.kpi('Ledger cost', w.money(totals.get('ledger_cost'), currency),
              hint='Class 2 balances in the general ledger', icon='menu_book', tone='violet')
        w.kpi('Register depreciation', w.money(totals.get('register_accum'), currency),
              hint='Accumulated per the register', icon='trending_down', tone='slate')
        w.kpi('Ledger depreciation', w.money(totals.get('ledger_accum'), currency),
              hint='Accumulated per the ledger', icon='trending_down', tone='slate')
        w.kpi('Reconciled', 'Yes' if reconciled else 'No',
              hint='Every account agrees to the cent' if reconciled
                   else 'At least one account disagrees',
              icon='verified' if reconciled else 'report_problem',
              tone='teal' if reconciled else 'rose')

    rows = data.get('rows') or []
    body = w.section(f"Register to ledger, as at {w.as_date(state['asOf'])}",
                     'The acceptance test the SRS asks for: the register and the ledger must '
                     'agree account by account')
    with body:
        for row in rows:
            row['id'] = f"{row['account']}-{row['kind']}"
            row['kind_label'] = 'Gross cost' if row['kind'] == 'gross' else 'Accumulated depreciation'
            row['agrees'] = 'Agrees' if float(row.get('variance') or 0) == 0 else 'Variance'
            row['categories_label'] = ', '.join(row.get('categories') or [])
        table = w.table([
            w.column('account', 'Account', classes='lc-num font-semibold'),
            w.column('account_name', 'Account name'),
            w.column('kind_label', 'Column'),
            w.column('categories_label', 'Categories'),
            w.int_column('asset_count', 'Assets'),
            w.money_column('register', f'Register ({currency})', currency=currency),
            w.money_column('ledger', f'Ledger ({currency})', currency=currency),
            w.money_column('variance', f'Variance ({currency})', currency=currency),
            w.column('agrees', 'Result'),
        ], rows, pagination=0, height='52vh')
        w.add_badge_slot(table, 'agrees', tones={'agrees': 'emerald', 'variance': 'rose'})

    export_state.update(
        title=f"Register to ledger reconciliation as at {w.as_date(state['asOf'])}",
        headers=['Account', 'Name', 'Column', 'Register', 'Ledger', 'Variance'],
        rows=[[r['account'], r.get('account_name'), r['kind_label'], r.get('register'),
               r.get('ledger'), r.get('variance')] for r in rows])


async def _depreciation_report(state: dict, currency: str, export_state: dict) -> None:
    data = await api.safe(api.get('/reports/depreciation',
                                  **{'from': state['from'], 'to': state['to']}))
    if not data:
        return
    totals = data.get('totals') or {}

    with w.kpi_row():
        w.kpi('OHADA charge', w.money(totals.get('ohada'), currency),
              hint='Statutory depreciation in the window', icon='trending_down', tone='sky')
        w.kpi('IFRS charge', w.money(totals.get('ifrs'), currency),
              hint='IFRS depreciation in the window', icon='balance', tone='violet')
        w.kpi('Difference', w.money(totals.get('delta'), currency),
              hint='Posted as IFRS-only entries', icon='call_split',
              tone='amber' if float(totals.get('delta') or 0) else 'slate')

    left, right = w.two_column()
    with left:
        by_period = data.get('by_period') or []
        body = w.section('By period', 'Posted runs only')
        with body:
            if not by_period:
                w.empty_state('No depreciation has been posted in this window.',
                              icon='trending_down')
            else:
                w.statement_table(
                    ['Period', 'Assets', f'OHADA ({currency})', f'IFRS ({currency})', 'Difference'],
                    [{'label': row['period_code'],
                      'values': [row.get('asset_count'), row.get('ohada'), row.get('ifrs'),
                                 row.get('delta')]} for row in by_period]
                    + [{'label': 'Total', 'kind': 'total',
                        'values': [None, totals.get('ohada'), totals.get('ifrs'),
                                   totals.get('delta')]}],
                    currency)
    with right:
        by_category = data.get('by_category') or []
        body = w.section('By category')
        with body:
            if not by_category:
                ui.label('Nothing to show.').classes('text-xs text-slate-500')
            else:
                w.statement_table(
                    ['Category', f'OHADA ({currency})', f'IFRS ({currency})'],
                    [{'code': row.get('category_code'), 'label': row.get('category_name'),
                      'values': [row.get('ohada'), row.get('ifrs')]} for row in by_category],
                    currency)

    export_state.update(
        title=f"Depreciation charge {w.as_date(state['from'])} to {w.as_date(state['to'])}",
        headers=['Period', 'Assets', 'OHADA', 'IFRS', 'Difference'],
        rows=[[r['period_code'], r.get('asset_count'), r.get('ohada'), r.get('ifrs'), r.get('delta')]
              for r in data.get('by_period') or []])


async def _compliance_report(state: dict, export_state: dict) -> None:
    data = await api.safe(api.get('/reports/compliance', asOf=state['asOf']))
    if not data:
        return

    with w.kpi_row():
        w.kpi('Overall compliance', w.percent(data.get('overall_compliance_pct')),
              hint='Devices holding a current calibration certificate',
              icon='verified_user',
              tone='teal' if float(data.get('overall_compliance_pct') or 0) >= 95 else 'amber')
        w.kpi('Past a due date', w.number(len(data.get('overdue') or [])),
              hint='Calibration, safety test or preventive service', icon='event_busy',
              tone='rose' if data.get('overdue') else 'teal')
        w.kpi('Due within 60 days', w.number(len(data.get('due_soon') or [])),
              hint='Plan the work now', icon='schedule', tone='amber')
        w.kpi('Restricted from use', w.number(data.get('restricted')),
              hint='Cannot be used on patients', icon='gpp_bad',
              tone='rose' if data.get('restricted') else 'teal')

    coverage = data.get('calibration_coverage') or []
    body = w.section('Calibration coverage by domain',
                     f"As at {w.as_date(data.get('as_of'))} — the position a regulator asks for")
    with body:
        if not coverage:
            w.empty_state('No category requires calibration.', icon='straighten')
        else:
            for row in coverage:
                row['id'] = row['domain']
                row['domain_label'] = w.humanise(row['domain'])
                row['compliance'] = w.percent(row.get('compliance_pct'))
            w.table([w.column('domain_label', 'Domain'),
                     w.int_column('assets_requiring', 'Devices requiring calibration'),
                     w.int_column('in_date', 'In date'),
                     w.int_column('overdue', 'Overdue'),
                     w.int_column('never_calibrated', 'Never calibrated'),
                     w.column('compliance', 'Compliance', align='right', classes='lc-num')],
                    coverage, pagination=0)

    overdue = data.get('overdue') or []
    body = w.section(f'{w.number(len(overdue))} assets past a due date',
                     'Any lapsed calibration, safety test or preventive service — wider than '
                     'the calibration coverage above')
    with body:
        if not overdue:
            w.empty_state('Nothing is overdue.', icon='verified')
        else:
            for row in overdue:
                row['asset'] = f"{row.get('asset_tag')} — {row.get('name')}"
                row['state'] = 'Restricted' if row.get('use_restricted') else 'Overdue'
            table = w.table([w.column('asset', 'Asset'), w.column('domain', 'Domain'),
                             w.date_column('next_calibration_date', 'Calibration due'),
                             w.date_column('next_safety_test_date', 'Safety test due'),
                             w.date_column('next_pm_date', 'Service due'),
                             w.column('state', 'State'),
                             w.column('restriction_reason', 'Restriction')],
                            overdue, pagination=0, height='42vh',
                            on_row_click=lambda row: ui.navigate.to(f"/assets/{row['id']}"))
            w.add_badge_slot(table, 'state', tones={'restricted': 'rose', 'overdue': 'amber'})

    contracts = data.get('contracts') or []
    if contracts:
        body = w.section('Service contracts', 'Cover that is running out')
        with body:
            table = w.table([w.column('contract_no', 'Contract', classes='lc-num'),
                             w.column('title', 'Cover'),
                             w.column('supplier_name', 'Supplier'),
                             w.int_column('asset_count', 'Assets'),
                             w.date_column('end_date', 'Ends'),
                             w.int_column('days_to_expiry', 'Days left'),
                             w.column('status', 'Status')],
                            contracts, pagination=10)
            w.add_badge_slot(table, 'status')

    export_state.update(
        title=f"Compliance position as at {w.as_date(data.get('as_of'))}",
        headers=['Domain', 'Requiring calibration', 'In date', 'Overdue', 'Never calibrated',
                 'Compliance %'],
        rows=[[w.humanise(r['domain']), r.get('assets_requiring'), r.get('in_date'),
               r.get('overdue'), r.get('never_calibrated'), r.get('compliance_pct')]
              for r in coverage])


async def _tco_report(state: dict, currency: str, export_state: dict) -> None:
    data = await api.safe(api.get('/reports/cost-of-ownership',
                                  **{'from': state['from'], 'to': state['to']}))
    if not data:
        return
    rows = data.get('rows') or []
    candidates = data.get('replacement_candidates') or []

    with w.kpi_row():
        w.kpi('Assets analysed', w.number(data.get('total')),
              hint='With maintenance history in the window', icon='inventory_2', tone='slate')
        w.kpi('Replacement candidates', w.number(len(candidates)),
              hint='Lifetime repairs are a third or more of cost',
              icon='autorenew', tone='rose' if candidates else 'teal')
        if rows:
            availability = sum(float(r.get('availability_pct') or 0) for r in rows) / len(rows)
            w.kpi('Average availability', w.percent(availability),
                  hint='Uptime across the window', icon='timelapse',
                  tone='teal' if availability >= 95 else 'amber')

    if candidates:
        body = w.section('Replace rather than repair',
                         'Equipment whose lifetime maintenance has reached a third of what it '
                         'cost — the point at which repair stops being economic')
        with body:
            for row in candidates:
                row['id'] = row['asset_tag']
                row['ratio'] = w.percent(row.get('maintenance_to_cost_pct'))
            table = w.table([w.column('asset_tag', 'Tag', classes='lc-num'),
                             w.column('name', 'Asset'),
                             w.column('criticality', 'Criticality'),
                             w.money_column('ohada_cost', f'Cost ({currency})', currency=currency),
                             w.money_column('total_maintenance_cost', f'Maintenance ({currency})',
                                            currency=currency),
                             w.column('ratio', 'Maintenance to cost', align='right',
                                      classes='lc-num'),
                             w.int_column('failures', 'Failures')],
                            candidates, pagination=10)
            w.add_badge_slot(table, 'criticality')

    body = w.section('Reliability and cost of ownership',
                     f"{w.as_date(state['from'])} to {w.as_date(state['to'])} · availability, "
                     'mean time between failures and mean time to repair')
    with body:
        if not rows:
            w.empty_state('No maintenance activity in this window.', icon='build')
            return
        for row in rows:
            row['availability'] = w.percent(row.get('availability_pct'))
            row['ratio'] = w.percent(row.get('maintenance_to_cost_pct'))
        w.table([w.column('asset_tag', 'Tag', classes='lc-num'),
                 w.column('name', 'Asset'),
                 w.column('domain', 'Domain'),
                 w.column('cost_center_name', 'Cost centre'),
                 w.money_column('ohada_cost', f'Cost ({currency})', currency=currency),
                 w.money_column('total_maintenance_cost', f'Maintenance ({currency})',
                                currency=currency),
                 w.column('ratio', 'Maint. to cost', align='right', classes='lc-num'),
                 w.int_column('work_orders', 'Work orders'),
                 w.int_column('failures', 'Failures'),
                 w.column('mtbf_hours', 'MTBF (h)', align='right', classes='lc-num'),
                 w.column('mttr_hours', 'MTTR (h)', align='right', classes='lc-num'),
                 w.column('availability', 'Availability', align='right', classes='lc-num')],
                [{**row, 'id': row['asset_tag']} for row in rows], pagination=0, height='52vh')

    export_state.update(
        title=f"Cost of ownership {w.as_date(state['from'])} to {w.as_date(state['to'])}",
        headers=['Tag', 'Asset', 'Domain', 'Cost', 'Maintenance', 'Maint. to cost %',
                 'Work orders', 'Failures', 'Availability %'],
        rows=[[r['asset_tag'], r['name'], r.get('domain'), r.get('ohada_cost'),
               r.get('total_maintenance_cost'), r.get('maintenance_to_cost_pct'),
               r.get('work_orders'), r.get('failures'), r.get('availability_pct')]
              for r in rows])


# --- exports ----------------------------------------------------------------


def _export_excel(export_state: dict, currency: str) -> None:
    """A real workbook, not a CSV renamed: figures stay numeric and formatted."""
    if not export_state.get('rows'):
        ui.notify('Nothing on screen to export yet.', type='warning', position='top')
        return

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    book = Workbook()
    sheet = book.active
    sheet.title = 'Statement'

    title_font = Font(name='Calibri', size=14, bold=True, color='0F172A')
    head_font = Font(name='Calibri', size=10, bold=True, color='FFFFFF')
    head_fill = PatternFill('solid', fgColor='0F766E')
    thin = Side(style='thin', color='CBD5E1')

    sheet['A1'] = export_state['title']
    sheet['A1'].font = title_font
    sheet['A2'] = f"Amounts in {currency} · generated {datetime.now():%d %b %Y %H:%M}"
    sheet['A2'].font = Font(name='Calibri', size=9, italic=True, color='64748B')

    headers = export_state.get('headers') or []
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=4, column=column, value=header)
        cell.font = head_font
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    fmt = '#,##0' if currency.upper() in ('XAF', 'XOF') else '#,##0.00'
    for r, row in enumerate(export_state['rows'], start=5):
        for c, value in enumerate(row, start=1):
            cell = sheet.cell(row=r, column=c)
            if isinstance(value, (int, float)):
                cell.value = value
                cell.number_format = fmt
                cell.alignment = Alignment(horizontal='right')
            else:
                cell.value = '' if value is None else str(value)
            cell.border = Border(bottom=thin)
            cell.font = Font(name='Calibri', size=10)

    widths = [max(12, min(48, max((len(str(row[i])) if i < len(row) and row[i] is not None else 0)
                                  for row in [headers, *export_state['rows']]) + 3))
              for i in range(len(headers))]
    for i, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(row=4, column=i).column_letter].width = width
    sheet.freeze_panes = 'A5'

    buffer = io.BytesIO()
    book.save(buffer)
    ui.download(buffer.getvalue(), f"{_slug(export_state['title'])}.xlsx")


def _export_pdf(export_state: dict, currency: str, org: dict) -> None:
    if not export_state.get('rows'):
        ui.notify('Nothing on screen to export yet.', type='warning', position='top')
        return

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    headers = export_state.get('headers') or []
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4) if len(headers) > 4 else A4,
        leftMargin=14 * mm, rightMargin=14 * mm, topMargin=14 * mm, bottomMargin=14 * mm,
        title=export_state['title'],
    )
    styles = getSampleStyleSheet()
    heading = ParagraphStyle('h', parent=styles['Title'], fontSize=14, alignment=0,
                             textColor=colors.HexColor('#0f172a'), spaceAfter=2)
    sub = ParagraphStyle('s', parent=styles['Normal'], fontSize=8,
                         textColor=colors.HexColor('#64748b'))
    cell = ParagraphStyle('c', parent=styles['Normal'], fontSize=7.5, leading=10)

    story = [
        Paragraph(org.get('name', 'LedgerCore Health'), heading),
        Paragraph(export_state['title'], ParagraphStyle('t', parent=styles['Normal'], fontSize=10,
                                                        spaceAfter=2)),
        Paragraph(f"Amounts in {currency} · generated {datetime.now():%d %B %Y at %H:%M}", sub),
        Spacer(1, 6 * mm),
    ]

    decimals = 0 if currency.upper() in ('XAF', 'XOF') else 2
    body = [[Paragraph(f'<b>{h}</b>', cell) for h in headers]]
    for row in export_state['rows']:
        rendered = []
        for value in row:
            if isinstance(value, (int, float)):
                text = f'{abs(value):,.{decimals}f}'.replace(',', ' ')
                rendered.append(f'({text})' if value < 0 else text)
            else:
                rendered.append('' if value is None else str(value))
        body.append([Paragraph(v, cell) for v in rendered])

    table = Table(body, repeatRows=1, hAlign='LEFT')
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0f766e')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (1, 1), (-1, -1), 'RIGHT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('LINEBELOW', (0, 0), (-1, -1), 0.4, colors.HexColor('#e2e8f0')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')]),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    story.append(table)
    doc.build(story)
    ui.download(buffer.getvalue(), f"{_slug(export_state['title'])}.pdf")


def _slug(text: str) -> str:
    keep = [c.lower() if c.isalnum() else '-' for c in text]
    slug = ''.join(keep)
    while '--' in slug:
        slug = slug.replace('--', '-')
    return slug.strip('-')[:70] or 'statement'
